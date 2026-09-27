import json
import logging

import pytest
import requests
from openpyxl import load_workbook

import main as cli
from budget import BudgetExceeded, EntityBudgetManager
from cache_manager import CacheManager
from collector import Collector, plan
from config import validate_config
from exporter import export
from logger import SecretFilter
from models import SchemaError, location, reviews
from storage import StateError, output_lock, read_json, write_json
from tripadvisor_client import APIError, TripadvisorClient, retry_delay


def config(maximum=5, second=False):
    # Synthetic coordinates used only by offline fakes, never real search instructions.
    regions = [{"name": "Fixture A", "lat": 0.0, "lon": 0.0, "radius_km": 1.0}]
    if second:
        regions.append({**regions[0], "name": "Fixture B"})
    return {"city": "부산", "regions": regions,
            "categories": ["RESTAURANT"], "max_locations": maximum}


def raw_place(identity):
    return {"id": identity, "names": [{"value": f"장소 {identity}", "language": "ko"}],
            "addresses": [{"formatted": "검증용 주소"}],
            "coordinates": {"latitude": 0, "longitude": 0},
            "traveler_ratings": {"overall": {"rating": 4.5, "count": 12}},
            "photos": {"total_count": 99}}


def raw_review(identity):
    return {"id": identity, "rating": 5, "title": [{"value": "좋아요", "language": "ko"}],
            "text": [{"value": "한글 후기", "language": "ko", "primary": True}],
            "publish_ts": "2026-09-01T10:00:00Z", "travel_date": "2026-08",
            "trip_type": "FAMILY", "url": "https://www.tripadvisor.com/",
            "photos": [{"original_size_url": "NEVER_SAVE_PHOTO"}],
            "user": {"avatar": "NEVER_SAVE_PHOTO"}}


class Response:
    def __init__(self, data=None, status=200, headers=None):
        self.data, self.status_code, self.headers = data, status, headers or {}

    def json(self):
        return self.data

    def close(self):
        pass


class Session:
    def __init__(self, responses):
        self.responses, self.calls, self.headers = list(responses), [], {}

    def get(self, url, **kwargs):
        self.calls.append((url, kwargs))
        response = self.responses.pop(0)
        if isinstance(response, Exception):
            raise response
        return response

    def close(self):
        pass


class FakeClient:
    def __init__(self, budget, fail=None, interrupt=None, empty=False):
        self.budget, self.calls = budget, []
        self.fail, self.interrupt, self.empty = fail, interrupt, empty

    def nearby(self, region, category, page, city):
        self.budget.reserve(5)
        self.calls.append(("search", page))
        items = [] if self.empty else [
            {"location": raw_place(i)} for i in range((page - 1) * 5 + 1, page * 5 + 1)]
        return {"data": items, "pagination": {"total_pages": 2}}

    def details(self, identity):
        self.budget.reserve(1)
        self.calls.append(("detail", identity))
        return raw_place(identity)

    def reviews(self, identity, name=None):
        self.budget.reserve(1)
        self.calls.append(("review", identity))
        if identity == self.interrupt:
            raise KeyboardInterrupt
        if identity == self.fail:
            raise APIError("offline fixture error")
        return {"data": [raw_review(int(identity) * 10 + i) for i in range(3)]}


def run_fake(root, cfg, refresh=False, **kwargs):
    budget = EntityBudgetManager(root / "entity_usage.json", 800)
    cache = CacheManager(root / "cache")
    client = FakeClient(budget, **kwargs)
    result = Collector(cfg, cache, budget, client, root, refresh).run()
    return budget, cache, client, result


def test_budget_accumulates_100_to_300(tmp_path):
    path = tmp_path / "entity_usage.json"
    first = EntityBudgetManager(path, 100)
    first.reserve(54)
    second = EntityBudgetManager(path, 300)
    assert second.estimated_used == 54 and second.remaining == 246
    second.reserve(246)
    with pytest.raises(BudgetExceeded):
        second.reserve(1)
    assert read_json(path)["estimated_used"] == 300


def test_budget_lower_limit_keeps_usage(tmp_path):
    path = tmp_path / "entity_usage.json"
    EntityBudgetManager(path, 300).reserve(150)
    budget = EntityBudgetManager(path, 100)
    assert budget.remaining == 0 and budget.estimated_used == 150
    with pytest.raises(BudgetExceeded):
        budget.reserve(1)


@pytest.mark.parametrize("value", [None, {}, {"estimated_used": -1}, {"estimated_used": True}])
def test_invalid_ledger_fails_closed(tmp_path, value):
    path = tmp_path / "entity_usage.json"
    write_json(path, value)
    with pytest.raises(StateError):
        EntityBudgetManager(path)


def test_missing_ledger_with_cache_fails_closed(tmp_path):
    write_json(tmp_path / "cache/reviews/1.json", {})
    with pytest.raises(StateError):
        EntityBudgetManager(tmp_path / "entity_usage.json")


def test_lock_blocks_concurrent_run(tmp_path):
    with output_lock(tmp_path):
        with pytest.raises(StateError):
            with output_lock(tmp_path):
                pytest.fail("second lock acquired")
    with output_lock(tmp_path):
        pass


def test_cache_miss_empty_hit_and_corruption(tmp_path):
    cache = CacheManager(tmp_path)
    assert cache.get("reviews", "123") is None
    cache.put("reviews", "123", {"reviews": [], "fetched_at": "now"})
    assert cache.get("reviews", "123")["reviews"] == []
    cache.path("reviews", "123").write_text("{", encoding="utf-8")
    with pytest.raises(StateError):
        cache.get("reviews", "123")


def test_duplicate_resume_and_increase_limit(tmp_path):
    first, cache, client, result = run_fake(tmp_path, config(second=True))
    assert first.estimated_used == 20  # Two searches, five details, five reviews.
    assert len([c for c in client.calls if c[0] == "review"]) == 5
    places = read_json(tmp_path / "reviews.json")
    assert len(places) == 5 and len(places[0]["regions"]) == 2
    assert all(p["collected_review_count"] == 3 for p in places)
    second, _, client, _ = run_fake(tmp_path, config(second=True))
    assert not client.calls and second.estimated_used == 20
    estimate = plan(config(10, second=True), cache, second)
    third, _, client, _ = run_fake(tmp_path, config(10, second=True))
    assert third.estimated_used == 40
    assert {c[1] for c in client.calls if c[0] == "review"} == {"6", "7", "8", "9", "10"}
    assert third.estimated_used - second.estimated_used <= estimate["additional_entities_without_retries"]


def test_refresh_deduplicates_and_accumulates(tmp_path):
    run_fake(tmp_path, config(second=True))
    budget, _, client, _ = run_fake(tmp_path, config(second=True), refresh=True)
    assert budget.estimated_used == 40
    assert len([c for c in client.calls if c[0] == "review"]) == 5


def test_review_error_not_cached_and_only_failed_place_resumes(tmp_path):
    _, cache, _, result = run_fake(tmp_path, config(), fail="2")
    assert result["failures"] == 1
    assert cache.get("reviews", "2") is None
    _, _, client, _ = run_fake(tmp_path, config())
    assert client.calls == [("review", "2")]


def test_keyboard_interrupt_resumes_from_pending(tmp_path):
    _, cache, _, result = run_fake(tmp_path, config(), interrupt="2")
    assert result["stopped"] == "interrupted"
    assert cache.get("reviews", "1") is not None
    _, _, client, _ = run_fake(tmp_path, config())
    assert ("review", "1") not in client.calls
    assert {c[1] for c in client.calls if c[0] == "review"} == {"2", "3", "4", "5"}


def test_budget_stop_exports_partial_and_resumes(tmp_path):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json", 8)
    cache = CacheManager(tmp_path / "cache")
    result = Collector(config(), cache, budget, FakeClient(budget), tmp_path).run()
    assert result["stopped"] == "budget" and budget.estimated_used == 8
    assert (tmp_path / "reviews.xlsx").exists()
    _, _, client, result = run_fake(tmp_path, config())
    assert ("review", "1") not in client.calls
    assert result["locations"] == 5


def test_empty_search_is_cached(tmp_path):
    budget, _, _, result = run_fake(tmp_path, config(), empty=True)
    assert result["locations"] == 0 and budget.estimated_used == 5
    _, _, client, _ = run_fake(tmp_path, config())
    assert not client.calls


def test_official_schema_optional_fields_and_no_photos():
    place = location(raw_place(1), "RESTAURANT")
    assert place["rating"] == 4.5 and place["review_count"] == 12
    assert location({"id": 2}, "HOTEL")["latitude"] is None
    result = reviews({"data": [raw_review(1), {}, raw_review(3), raw_review(4)]})
    assert len(result) == 3 and result[0]["published_date"] == "2026-09-01T10:00:00Z"
    assert result[1]["review_id"] is None
    assert "photos" not in json.dumps(result) and "NEVER_SAVE" not in json.dumps(result)
    with pytest.raises(SchemaError):
        reviews({"error": "wrong shape"})


def test_json_excel_output_and_formula_safety(tmp_path):
    place = {**location(raw_place(1), "HOTEL"), "regions": ["지역"],
             "collected_review_count": 1, "reviews": reviews({"data": [raw_review(1)]})}
    place["reviews"][0]["text"] = '=HYPERLINK("malicious")\x01'
    assert export(tmp_path, [place])
    assert read_json(tmp_path / "reviews.json")[0]["reviews"][0]["text"].endswith("\x01")
    book = load_workbook(tmp_path / "reviews.xlsx")
    assert book.sheetnames == ["Locations", "Reviews"]
    for sheet in book:
        assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref
    assert book["Reviews"]["I2"].data_type == "s"
    assert book["Reviews"]["I2"].alignment.wrap_text
    assert "\x01" not in book["Reviews"]["I2"].value
    book.close()


def test_http_budget_prevents_network_and_reserves_before_attempt(tmp_path):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json", 1)
    session = Session([Response({"data": []})])
    client = TripadvisorClient("secret", budget, session=session, sleep=lambda n: None)
    client.reviews("1")
    assert read_json(budget.path)["estimated_used"] == 1
    with pytest.raises(BudgetExceeded):
        client.reviews("2")
    assert len(session.calls) == 1
    assert session.calls[0][1]["allow_redirects"] is False


def test_failed_ledger_write_sends_nothing(tmp_path, monkeypatch):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json")
    session = Session([])
    client = TripadvisorClient("secret", budget, session=session)
    def fail(*args):
        raise OSError("disk full")
    monkeypatch.setattr(budget, "save", fail)
    with pytest.raises(OSError):
        client.reviews("1")
    assert not session.calls


def test_search_reserves_page_cost_and_rate_limits(tmp_path):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json")
    session = Session([Response({"data": []}), Response({"data": []})])
    sleeps = []
    client = TripadvisorClient("secret", budget, session=session, sleep=sleeps.append, clock=lambda: 0)
    for page in (1, 2):
        client.nearby(config()["regions"][0], "HOTEL", page, "fixture")
    assert budget.estimated_used == 10
    assert 1.1 in sleeps
    assert all(c[1]["params"]["size"] == 5 for c in session.calls)


def test_retries_charge_every_attempt_and_use_retry_after(tmp_path):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json")
    session = Session([Response(status=429, headers={"Retry-After": "4"}),
                       requests.Timeout("secret must never be logged"), Response({"data": []})])
    sleeps = []
    client = TripadvisorClient("secret", budget, session=session, sleep=sleeps.append)
    assert client.reviews("1") == {"data": []}
    assert budget.estimated_used == 3 and 4 in sleeps
    assert retry_delay("Wed, 01 Jan 2020 00:00:00 GMT", 1) == 2


def test_429_retries_are_bounded(tmp_path):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json")
    session = Session([Response(status=429)] * 3)
    client = TripadvisorClient("secret", budget, session=session, sleep=lambda n: None)
    with pytest.raises(APIError):
        client.reviews("1")
    assert len(session.calls) == 3 and budget.estimated_used == 3


@pytest.mark.parametrize("endpoint", ["/locations", "/locations/1/photos", "/photos", "https://evil"])
def test_forbidden_endpoints(endpoint, tmp_path):
    budget = EntityBudgetManager(tmp_path / "entity_usage.json")
    session = Session([])
    client = TripadvisorClient("secret", budget, session=session)
    with pytest.raises(ValueError):
        client.get(endpoint, {})
    assert not session.calls and budget.estimated_used == 0


def test_secret_redaction():
    record = logging.LogRecord("collector", logging.ERROR, "", 1, "error %s", ("SECRET",), None)
    SecretFilter("SECRET").filter(record)
    assert "SECRET" not in record.getMessage()


@pytest.mark.parametrize("value", [float("nan"), float("inf"), 91, True])
def test_invalid_coordinates_rejected(value):
    cfg = config()
    cfg["regions"][0]["lat"] = value
    with pytest.raises(ValueError):
        validate_config(cfg)


def test_dry_run_never_constructs_client_or_changes_ledger(tmp_path, monkeypatch, capsys):
    cfg_path = tmp_path / "config.json"
    write_json(cfg_path, config())
    output = tmp_path / "output"
    EntityBudgetManager(output / "entity_usage.json", 100).reserve(54)
    before = (output / "entity_usage.json").read_bytes()
    monkeypatch.setattr(cli, "OUTPUT_DIR", output)
    monkeypatch.setattr(cli, "BASE_DIR", tmp_path)
    def blocked(*args, **kwargs):
        pytest.fail("dry run must not construct a client")
    monkeypatch.setattr(cli, "TripadvisorClient", blocked)
    assert cli.main(["--dry-run", "--config", str(cfg_path)]) == 0
    assert before == (output / "entity_usage.json").read_bytes()
    assert "DRY RUN" in capsys.readouterr().out


def test_refresh_requires_both_confirmations(tmp_path, monkeypatch):
    from config import Settings
    cfg_path = tmp_path / "config.json"
    write_json(cfg_path, config())
    monkeypatch.setattr(cli, "OUTPUT_DIR", tmp_path / "output")
    monkeypatch.setattr(cli, "BASE_DIR", tmp_path)
    monkeypatch.setattr(cli.Settings, "load", lambda: Settings("offline-key", 800))
    def blocked(*args, **kwargs):
        pytest.fail("declined run must not construct client")
    monkeypatch.setattr(cli, "TripadvisorClient", blocked)
    for answers in (["n"], ["y", "n"]):
        iterator = iter(answers)
        monkeypatch.setattr("builtins.input", lambda prompt: next(iterator))
        assert cli.main(["--refresh", "--config", str(cfg_path)]) == 0
    assert not (tmp_path / "output/entity_usage.json").exists()
