import logging
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from storage import write_json

LOCATION_COLUMNS = ["regions", "category", "location_id", "name", "address", "latitude",
                    "longitude", "rating", "review_count", "collected_review_count"]
REVIEW_COLUMNS = ["regions", "category", "location_id", "location_name", "location_rating",
                  "review_id", "review_rating", "title", "text", "trip_type", "travel_date",
                  "published_date", "language", "review_url"]
log = logging.getLogger("collector")


def export(root: Path, locations):
    root = Path(root)
    data = sorted(locations, key=lambda x: x["location_id"])
    # JSON is the lossless copy, even when an Excel file is open/locked.
    write_json(root / "reviews.json", data)
    book = Workbook()
    book.remove(book.active)
    review_rows = []
    for place in data:
        for review in place.get("reviews", []):
            review_rows.append({**review, "regions": place["regions"],
                                "category": place["category"],
                                "location_id": place["location_id"],
                                "location_name": place["name"],
                                "location_rating": place["rating"],
                                "review_rating": review["rating"]})
    for name, columns, rows in (("Locations", LOCATION_COLUMNS, data),
                                ("Reviews", REVIEW_COLUMNS, review_rows)):
        sheet = book.create_sheet(name)
        sheet.append(columns)
        for row in rows:
            values = []
            for column in columns:
                value = row.get(column)
                if isinstance(value, list):
                    value = ", ".join(value)
                if isinstance(value, str):
                    value = ILLEGAL_CHARACTERS_RE.sub("", value)
                    if len(value) > 32767:
                        log.warning("Excel cell truncated; complete text remains in reviews.json")
                    value = value[:32767]
                values.append(value)
            sheet.append(values)
            for cell in sheet[sheet.max_row]:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Untrusted review text must never become a formula.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176B59")
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[1].height = 32
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for idx, column in enumerate(columns, 1):
            width = 70 if column == "text" else 40 if column in (
                "address", "name", "location_name", "title", "review_url") else 24
            sheet.column_dimensions[get_column_letter(idx)].width = width
        for idx in range(2, sheet.max_row + 1):
            sheet.row_dimensions[idx].height = 75 if name == "Reviews" else 36
    fd, name = tempfile.mkstemp(dir=root, suffix=".xlsx")
    os.close(fd)
    try:
        book.save(name)
        os.replace(name, root / "reviews.xlsx")
    except PermissionError:
        log.error("reviews.xlsx가 열려 있거나 쓰기 권한이 없습니다. JSON/캐시는 저장되었습니다.")
        return False
    finally:
        book.close()
        if os.path.exists(name):
            os.unlink(name)
    return True
